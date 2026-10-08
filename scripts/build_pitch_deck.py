"""Build the Parali Se Paisa judge deck from fixed, reviewed research outputs."""
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/pitch/Parali_Se_Paisa_Greenovators2026.pptx"
FIG = ROOT / "reports/research/figures"
P = Presentation()
P.slide_width, P.slide_height = Inches(13.333), Inches(7.5)
BG=RGBColor(248,249,245); INK=RGBColor(27,40,33); GREEN=RGBColor(31,101,67); MUTED=RGBColor(91,105,95); PALE=RGBColor(230,239,230); GOLD=RGBColor(222,166,57); WHITE=RGBColor(255,255,255)

def txt(slide,x,y,w,h,value,size=16,color=INK,bold=False,align=PP_ALIGN.LEFT):
    sh=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=sh.text_frame; tf.word_wrap=True
    for i,line in enumerate(value.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align; p.space_after=Pt(5)
        r=p.add_run(); r.text=line; r.font.name="Aptos"; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return sh

def base(kicker,title,source,notes):
    s=P.slides.add_slide(P.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
    txt(s,.55,.28,12,.28,kicker.upper(),11,GREEN,True); txt(s,.55,.62,12.1,.75,title,27,INK,True)
    if source: txt(s,.55,7.02,11.7,.28,source,8,MUTED)
    txt(s,12.15,7.02,.6,.28,str(len(P.slides)),8,MUTED,False,PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text=notes
    return s

def card(s,x,y,w,h,heading,body,accent=GREEN):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=WHITE; sh.line.color.rgb=PALE
    sh.text_frame.clear(); sh.text_frame.word_wrap=True; sh.text_frame.margin_left=Inches(.16); sh.text_frame.margin_right=Inches(.14); sh.text_frame.margin_top=Inches(.12)
    p=sh.text_frame.paragraphs[0]; r=p.add_run(); r.text=heading; r.font.size=Pt(15); r.font.bold=True; r.font.color.rgb=accent
    p=sh.text_frame.add_paragraph(); p.space_before=Pt(8); r=p.add_run(); r.text=body; r.font.size=Pt(12); r.font.color.rgb=INK

def img(s,path,x,y,w,h):
    s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))

def img_fit(s,path,x,y,w,h):
    from PIL import Image
    with Image.open(path) as image:
        ratio=image.width/image.height
    width=min(w,h*ratio); height=width/ratio
    s.shapes.add_picture(str(path), Inches(x+(w-width)/2), Inches(y+(h-height)/2), width=Inches(width), height=Inches(height))

# 1 Problem
s=base("Problem","After harvest, straw needs a buyer and a baler","Source: Parali Se Paisa concept note; operational workflow shown later is a synthetic prototype.","Open with the practical coordination problem: residue has a short handling window; farmers, balers and buyers need a way to coordinate pickup. Our live evidence is satellite observation. The operating workflow is a prototype and uses synthetic records.")
txt(s,.7,1.8,6.1,2.15,"A narrow window. Many fields. Scattered machinery and demand.",26,INK,True)
card(s,7.2,1.7,2.55,2.2,"Farmers","Need a timely straw pickup option.")
card(s,9.95,1.7,2.55,2.2,"Operators","Need a clear dispatch queue.")
card(s,7.2,4.15,5.3,1.55,"Buyers","Need visible, traceable supply; this prototype does not establish actual demand.")

# 2 Solution
s=base("Solution","One evidence-to-action workflow","Source: application prototype; see README for real-data and synthetic-mode boundaries.","Explain the separation. Real mode is a read-only research view. The dispatch, buyer and evidence workflows are a separate synthetic prototype. No live service, verified registry, or incentive decision is claimed.")
for i,(h,b) in enumerate([("Observe","Sentinel-2 field histories"),("Prioritize","Normalized burn risk score, not probability"),("Coordinate","Demo baler/buyer workflow"),("Record","Prototype QR metadata check")]): card(s,.65+i*3.15,2.05,2.8,2.3,h,b)
txt(s,.75,5.3,11.8,.8,"REAL DATA  →  RULE-BASED INTERPRETATION  →  DEMO/SYNTHETIC OPERATIONS  →  PROTOTYPE RECORD",15,GREEN,True,PP_ALIGN.CENTER)

# 3 Real data
s=base("Real data","Muktsar map and one field history, 2023–2026","Source: fresh screenshots from isolated read-only real-data mode; 74 Sentinel-2 scenes, 12.6 million field-date observations; Fields of The World boundaries are model-derived.","Show the district candidate map, then the selected field's dense satellite history. Missing dates are not interpolated. The field is an example, and the burn tiers are rule-based candidates, not independently validated labels.")
txt(s,.7,1.43,12,.35,"74 Sentinel-2 scenes  ·  12.6 million field-date observations",13,GREEN,True)
img_fit(s,ROOT/"reports/pitch/research-real-map.png",.55,1.85,6.0,4.55)
img_fit(s,ROOT/"reports/pitch/field-history-real.png",6.8,1.85,6.0,4.55)
txt(s,.72,6.43,5.8,.32,"REAL DATA · rule-based candidate map",11,GREEN,True)
txt(s,6.95,6.43,5.8,.32,"REAL DATA · one field, stored observations",11,GREEN,True)

# 4 Real data + terminology
s=base("Real data · rule-based","Candidate area depends on the rule tier","Source: burn_candidate_summary.json and FINAL_JUDGE_EVIDENCE.md; full-revisit areas use non-overlapping 20 m pixels.","The loose candidate area is exploratory and is the cross-season series shown. Strict area is precision-first but cannot be read as a year-to-year trend because clean pre-event coverage differs with haze and cloud. FIRMS/VIIRS is corroborative thermal context, not ground truth.")
card(s,.75,1.7,5.7,3.9,"LOOSE_BURN_CANDIDATE","2023: 21,121 ha  |  2024: 10,566 ha  |  2025: 22,941 ha. Full-revisit, exploratory rule-based area.")
card(s,6.85,1.7,5.7,3.9,"STRICT_BURN_CANDIDATE","2023: 649 ha  |  2024: 1,022 ha  |  2025: 2,269 ha. Not a trend: haze/cloud changes clean pre-event coverage.",GOLD)
txt(s,.85,5.95,11.4,.5,"All outputs are RULE-BASED BURN CANDIDATES. Human review labels completed: 0.",13,INK,True)

# 5 Differentiator
s=base("Differentiator","Optical candidates and thermal context answer different questions","Source: 2025 strict-candidate / VIIRS matching analysis; defined ≤5-day evidence window; 500 m centroid radius.","In 2025, 8.8% of strict optical candidates had a matched VIIRS active-fire detection within the defined matching window. Say: About 9 in 10 strict optical burn candidates had no matched VIIRS active-fire detection within the defined matching window. Do not call this missed fires. Candidates are rule-based and VIIRS is corroborative thermal context.")
txt(s,.85,1.8,4.4,1.2,"8.8%",48,GREEN,True)
txt(s,.9,3.0,4.2,1.0,"of 2025 strict optical candidates had a matched VIIRS active-fire detection within the defined window.",16,INK)
card(s,5.25,1.8,7.15,1.85,"What the match means","About 9 in 10 strict optical burn candidates had no matched VIIRS active-fire detection within the defined matching window.")
card(s,5.25,4.0,7.15,1.5,"Interpretation","VIIRS is corroborative thermal context, not ground truth or a field-level confirmation.",GOLD)

# 6 Straw estimate
s=base("Action context","A reported straw figure needs its source","Source: The Tribune, as cited in reference 11 of Parali Se Paisa Concept Note; news-reported estimate, not an official statistic.","Attribute the 12.1 lakh tonne figure directly to The Tribune. It is a reported estimate and should not be represented as an official district statistic or as a measured result of our system.")
card(s,.8,1.85,5.5,2.5,"12.1 lakh tonnes","The Tribune reported this straw estimate. It is not an official statistic.",GOLD)
card(s,6.75,1.85,5.7,2.5,"From estimate to action","The product concept connects field evidence to straw pickup planning; actual supply and buyer demand still need field confirmation.")
txt(s,.9,5.0,11.4,.9,"Attribution stays with the number wherever it appears: The Tribune (reference 11 in the concept note).",14,MUTED)

# 7 Demo/prototype operations
s=base("Demo / synthetic","Dispatch is an operational prototype","Source: isolated local server, synthetic demo fixtures; route is a straight-line GEODESIC_PROXY only.","Clearly label every object in this segment DEMO/SYNTHETIC. Routing is a straight-line geodesic proxy and does not provide road travel time. Baler and buyer records are fixtures, not verified registries.")
img(s,ROOT/"reports/pitch/prototype-dispatch.png",.65,1.55,5.9,4.6)
img(s,ROOT/"reports/pitch/prototype-certificate-detail.png",6.8,1.55,5.9,4.6)
txt(s,.7,6.25,5.7,.45,"DEMO/SYNTHETIC · route proxy",11,GREEN,True)
txt(s,6.85,6.25,5.7,.45,"PROTOTYPE · QR metadata only",11,GREEN,True)

# 8 Limits
s=base("Limits","What the evidence can and cannot support","Source: FINAL_JUDGE_EVIDENCE.md; RESEARCH_SUMMARY.md; RESEARCH_METHODS.md.","We have not completed independent human labels. Strict and loose tiers are rule outputs. Risk is a normalized score, not a probability. Optical and thermal layers have different coverage limits. This is research decision support, not enforcement evidence.")
for i,(h,b) in enumerate([("No human validation yet","0 completed blinded labels"),("No strict trend","Pre-event clarity varies by season"),("No ground truth claim","Thermal matches are corroborative context"),("No probability claim","Risk is normalized score, not probability")]): card(s,.7+(i%2)*6.15,1.8+(i//2)*2.2,5.5,1.7,h,b,GOLD)

# 9 Action / close
s=base("Action","Pilot the evidence-to-pickup loop","Project sources: docs/RESEARCH_METHODS.md; docs/PITCH_AND_DEMO_GUIDE.md; synthetic prototype workflow.","Close by inviting the judges to inspect a real Muktsar field history, then show the separate synthetic dispatch and prototype QR metadata check. The next step is a bounded pilot with independent labels and real operator/buyer participation.")
txt(s,.8,1.9,11.6,1.15,"Real satellite evidence → clear limits → a prototype path from straw to value.",27,INK,True)
card(s,.9,4.35,3.55,1.55,"1 · Inspect","Muktsar real map and one field history")
card(s,4.9,4.35,3.55,1.55,"2 · Demonstrate","Synthetic dispatch and buyer workflow")
card(s,8.9,4.35,3.55,1.55,"3 · Pilot","Independent labels and local partners")
OUT.parent.mkdir(parents=True,exist_ok=True); P.save(OUT)
print(OUT)
